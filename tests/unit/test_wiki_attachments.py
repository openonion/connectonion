"""The terms are in the contract, not the mail that carried it."""

from pathlib import Path

from connectonion.wiki.attachments import extract_text


def test_a_pdf_is_read_as_text(tmp_path):
    from pypdf import PdfWriter
    pdf = tmp_path / "terms.pdf"
    writer = PdfWriter(); writer.add_blank_page(width=200, height=200); writer.write(str(pdf))
    assert isinstance(extract_text(pdf), str)      # empty page reads as empty, not a crash


def test_a_word_document_is_read_including_its_tables(tmp_path):
    from docx import Document
    doc = Document(); doc.add_paragraph("Fee: 7.5% of Net Booking Revenue")
    table = doc.add_table(rows=1, cols=2); table.rows[0].cells[0].text = "Term"; table.rows[0].cells[1].text = "90 days"
    path = tmp_path / "agreement.docx"; doc.save(str(path))
    text = extract_text(path)
    assert "7.5% of Net Booking Revenue" in text and "Term | 90 days" in text


def test_plain_text_and_html_are_read_and_markup_is_dropped(tmp_path):
    (tmp_path / "a.txt").write_text("hello   world\n", encoding="utf-8")
    (tmp_path / "b.html").write_text("<p>Signed <b>2026-08-07</b></p>", encoding="utf-8")
    assert extract_text(tmp_path / "a.txt") == "hello world"
    assert extract_text(tmp_path / "b.html") == "Signed 2026-08-07"


def test_an_unreadable_type_is_named_not_pretended(tmp_path):
    path = tmp_path / "photo.heic"; path.write_bytes(b"\x00\x01")
    text = extract_text(path)
    assert "photo.heic" in text and "not read" in text


def test_long_text_is_capped_and_says_so(tmp_path):
    (tmp_path / "big.txt").write_text("x " * 30_000, encoding="utf-8")
    text = extract_text(tmp_path / "big.txt", limit=100)
    assert len(text) < 200 and "more characters not shown" in text


def test_gather_reads_attachments_of_matched_mail_into_items(tmp_path, monkeypatch):
    from connectonion.wiki import investigate as inv
    monkeypatch.setattr("time.sleep", lambda s: None)

    class WithFiles:
        def my_addresses(self): return {"me@x.y"}
        def list_between(self, s, e, n):
            # window-aware, like a provider: one mail, in one week, not once per week
            rows = [{"id": "m1", "from": "Emma <szh526@gmail.com>", "to": ["me@x.y"], "cc": [],
                     "date": "2026-08-06T00:00:00+00:00", "subject": "contract v9"}]
            return [r for r in rows if s[:10] <= r["date"][:10] < e[:10]]
        def get_email_body(self, i): return "--- Email Body ---\nplease see attached"
        def download_attachments(self, email_id, out_dir):
            folder = Path(out_dir); folder.mkdir(parents=True, exist_ok=True)
            f = folder / "v9.txt"; f.write_text("Fee 7.5% of Net Booking Revenue", encoding="utf-8")
            return [str(f)]

    monkeypatch.setattr(inv, "datetime", __import__("datetime").datetime)
    items, coverage = inv.gather("Emma", ["szh526"], days=60, clients={"gmail": WithFiles()},
                                 subscriptions={}, attachments_dir=tmp_path / "att")
    kinds = [i["role"] for i in items]
    assert kinds == ["other", "attachment"]
    assert "7.5%" in items[1]["text"] and items[1]["source"].endswith(":v9.txt")
    assert any("1 attachments read" in c for c in coverage)
    assert (tmp_path / "att" / "gmail").is_dir()


def test_gmails_result_dict_yields_only_the_files_actually_saved():
    """Gmail reports per-file status; a failed row has no path and is not a file.
    Reading the wrong key would have returned [] for every Gmail mail, silently."""
    from connectonion.wiki.investigate import _saved_paths
    result = {"items": [{"id": "a", "filename": "deck.pdf", "status": "saved", "path": "/x/deck.pdf"},
                        {"id": "b", "filename": "bad.bin", "status": "failed", "error": "length mismatch"}],
              "complete": False, "decoded_bytes": 10}
    assert _saved_paths(result) == ["/x/deck.pdf"]
    assert _saved_paths(["/y/a.docx", "/y/b.txt"]) == ["/y/a.docx", "/y/b.txt"]   # Outlook's list
    assert _saved_paths(None) == []


def test_gather_creates_gmail_download_directory(tmp_path):
    from datetime import datetime, timedelta, timezone
    from connectonion.wiki.investigate import gather

    class Gmail:
        def my_addresses(self): return {"me@example.com"}
        def list_between(self, *args):
            return [{"id": "one", "from": "alice@example.com", "to": "me@example.com",
                     "date": (datetime.now(timezone.utc) - timedelta(hours=1)).isoformat(),
                     "subject": "Signed terms"}]
        def get_email_body(self, email_id): return "Please see attached"
        def download_attachments(self, email_id, directory, *, all_attachments=False):
            destination = Path(directory).resolve(strict=True)  # Gmail's actual contract.
            assert all_attachments
            path = destination / "terms.txt"
            path.write_text("Signed on 2026-09-14")
            return {"items": [{"status": "saved", "path": str(path)}], "complete": True}

    items, coverage = gather("Alice", ["alice@example.com"], days=1, clients={"gmail": Gmail()},
                             subscriptions={}, attachments_dir=tmp_path / "attachments")
    assert any(i["role"] == "attachment" and "Signed on" in i["text"] for i in items)
    assert not any("could not be saved" in line for line in coverage)
